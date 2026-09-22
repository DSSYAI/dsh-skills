# -*- coding: utf-8 -*-
"""可复用的 docx 生成器模板。运行在 uv venv：<venv>\Scripts\python.exe 本文件.py <输出路径>"""
import sys
from docx import Document
from docx.shared import Pt, RGBColor

def build(doc: Document):
    # —— 标题 / 章节 ——
    doc.add_heading("文档标题（覆盖这里）", 0)
    doc.add_heading("第一章　章节名", level=1)
    doc.add_heading("1.1　小节名", level=2)

    # —— 正文段落 ——
    doc.add_paragraph("正文内容。python-docx 默认正文样式是 Calibri 11。中英混排 OK。")

    # —— 加粗 / 强调 ——
    p = doc.add_paragraph("这段演示：")
    r = p.add_run("加粗文字")
    r.bold = True
    p.add_run("，和")
    r2 = p.add_run("斜体文字")
    r2.italic = True
    p.add_run("。")

    # 用样式 Intense Quote 做强调块
    q = doc.add_paragraph("要点提示……", style="Intense Quote")
    q.runs[0].bold = True

    # —— 列表 ——
    doc.add_paragraph("第一项", style="List Bullet")
    doc.add_paragraph("第二项", style="List Bullet")

    # —— 表格 ——（Table Grid 便于看格线）
    t = doc.add_table(rows=3, cols=2)
    t.style = "Table Grid"
    header = t.rows[0].cells
    header[0].text = "列A"; header[1].text = "列B"
    t.rows[1].cells[0].text = "甲"; t.rows[1].cells[1].text = "乙"
    t.rows[2].cells[0].text = "丙"; t.rows[2].cells[1].text = "丁"

def main(out: str):
    doc = Document()
    build(doc)
    doc.save(out)
    print("saved OK:", out)

if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "out.docx")